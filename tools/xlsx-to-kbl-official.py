#!/usr/bin/env python3
"""KBL 공식 자료 xlsx(프로필·시즌 기록·수상, 사용자 제공 2026-10-23) → CSV.
  python3 -I tools/xlsx-to-kbl-official.py <xlsx>
만드는 파일:
  data/players/births.csv          선수번호별 생년월일·신체·포지션·학교 (선수검색 + 추가 확인 명단)
  data/stats/official/<시즌>.csv   시즌·경기구분(정규/플레이오프/챔피언결정전)·소속팀별 선수 누적 기록 (팀별 시즌스탯 시트)
  data/history/awards.csv          역대 수상 경력(개인상·팀 성적)
수식 칸(출생연도·만 나이·MIN)은 읽지 않고 원값(생년월일 원문·출전시간(초))만 쓴다. 비어 있는 값은 추정하지 않고 빈 칸으로 둔다.
시즌 표기는 history와 같게 맞춘다(2026-2027 → 2026-27, 1999-2000은 그대로)."""
import sys, csv, datetime, os, re, collections
import openpyxl

xlsx = sys.argv[1]
wb = openpyxl.load_workbook(xlsx, data_only=False)
os.makedirs('data/stats/official', exist_ok=True)


def season_label(s):
    s = str(s).strip()
    m = re.fullmatch(r'(\d{4})-(\d{4})', s)
    if not m: return s  # '1997'
    a, b = int(m[1]), int(m[2])
    if b != a + 1: sys.exit(f'시즌 표기 이상: {s}')
    return s if b % 100 == 0 or a == 1999 else f'{a}-{str(b)[2:]}'


# ---------- 프로필 ----------
raw = {}
for r in wb['원본 및 안내'].iter_rows(min_row=11, values_only=True):
    if r[0] is None: continue
    raw[str(r[0])] = dict(nameEn=(r[2] or '').strip(), rawBirth=(str(r[9]).strip() if r[9] is not None else ''))

KIND = {'국내': 'domestic', '외국': 'foreign', '아시아쿼터': 'asia'}
num = lambda v: '' if v in (None, '') else (str(int(v)) if float(v) == int(float(v)) else str(v))
people = []
for sheet, group, first in (('선수 출생연도', 'main', 10), ('추가 확인 대상', 'extra', 8)):
    for r in wb[sheet].iter_rows(min_row=first, values_only=True):
        if r[0] is None: continue
        no = str(r[4])
        d = r[2]
        birth = d.strftime('%Y-%m-%d') if isinstance(d, (datetime.datetime, datetime.date)) else ''
        rb = raw.get(no, {}).get('rawBirth', '')
        if birth and rb and rb.replace('-', '') != birth.replace('-', ''):
            sys.exit(f'생년월일 불일치 {no} {r[0]}: {birth} vs 원문 {rb}')
        note = r[7] or ''
        if rb and not birth:  # 달력에 없는 날짜(원문 오류)는 값으로 쓰지 않고 원문을 남긴다
            note = (note + '; ' if note else '') + f'생년월일 원문 오류 {rb}'
        # 열 순서: J 포지션, K 키, L 몸무게, M 등번호, N 만 나이(수식, 읽지 않음), O 출신학교, P 프로필 응답시즌
        extra = list(r[9:17]) + [None] * (8 - len(r[9:17])) if group == 'main' else [None] * 8
        people.append(dict(playerNo=no, name=r[0], nameEn=raw.get(no, {}).get('nameEn', ''), kind=KIND.get(r[3], r[3]), team=(r[5] or ''), birth=birth,
                           status=r[6] or '', note=note, list=group,
                           position=extra[0] or '', heightCm=num(extra[1]), weightKg=num(extra[2]), number=str(extra[3] or ''), school=(extra[5] or '') if len(extra) > 5 else '',
                           profileSeason=extra[6] or ''))
with open('data/players/births.csv', 'w', newline='', encoding='utf-8') as f:
    f.write('# KBL 공식 프로필 API에서 수집한 프로필(사용자 제공 xlsx, 2026-10-23). 생년월일 빈 칸은 원자료에 값이 없는 경우(추정 안 함).\n'
            '# team은 기록상 구단 표기(현재/과거, 전체 경력 아님). list=extra는 공식 선수검색 명단에 없던 번호.\n'
            '# 키·몸무게·등번호·학교·포지션은 조회 시점(profileSeason 응답) 값이라 과거 시즌 당시 값과 다를 수 있다.\n')
    w = csv.DictWriter(f, fieldnames=list(people[0]), lineterminator='\n')
    w.writeheader(); w.writerows(people)
print(f'births.csv: {len(people)}행 (생년월일 {sum(1 for p in people if p["birth"])}, 키 {sum(1 for p in people if p["heightCm"])}, 학교 {sum(1 for p in people if p["school"])})')

# ---------- 시즌 기록 ----------
TYPE = {'정규시즌': 'regular', '플레이오프': 'playoff', '챔피언결정전': 'final'}
COLS = [('gp', 6), ('w', 7), ('l', 8), ('sec', 35), ('pts', 10), ('pm2', 11), ('pa2', 12), ('pm3', 14), ('pa3', 15), ('fgm', 17), ('fga', 18), ('ftm', 20), ('fta', 21),
        ('oreb', 23), ('dreb', 24), ('reb', 25), ('ast', 26), ('stl', 27), ('blk', 28), ('gd', 29), ('dk', 30), ('dka', 31), ('to', 32), ('pf', 33), ('pp', 34)]
by = collections.defaultdict(list)
for r in wb['팀별 시즌스탯'].iter_rows(min_row=8, values_only=True):
    if r[0] is None: continue
    row = dict(gameType=TYPE[r[1]], name=r[3], team=r[4], playerNo=str(r[5]), teamCode=str(r[36]))
    for k, i in COLS: row[k] = num(r[i])
    by[season_label(r[0])].append(row)
total = 0
for season, rows in by.items():
    head = ['gameType', 'name', 'team', 'teamCode', 'playerNo'] + [k for k, _ in COLS]
    with open(f'data/stats/official/{season}.csv', 'w', newline='', encoding='utf-8') as f:
        f.write('# KBL 공식 기록(사용자 제공 xlsx 팀별 시즌스탯 시트, 2026-10-23). 경기구분·소속팀별 한 줄(이적 선수는 팀마다 별도). sec=출전시간(초). 비율·평균은 계산 가능해 뺐다.\n')
        w = csv.DictWriter(f, fieldnames=head, lineterminator='\n')
        w.writeheader(); w.writerows(rows)
    total += len(rows)
print(f'data/stats/official: {len(by)}개 시즌 {total}행')

# ---------- 수상 ----------
one_line = lambda v: re.sub(r'\s*[\r\n]+\s*', ' / ', str(v)) if v else ''  # CSV를 한 줄에 한 행으로 유지
rows = []
for r in wb['수상 경력'].iter_rows(min_row=8, values_only=True):
    if r[0] is None: continue
    rows.append(dict(season=season_label(r[0]), holder=r[1] or '', category=r[2] or '', award=one_line(r[3]), name=one_line(r[4]), team=one_line(r[5]), playerNo=str(r[6] or ''), link=r[7] or '', raw=one_line(r[8])))
with open('data/history/awards.csv', 'w', newline='', encoding='utf-8') as f:
    f.write('# KBL 역대 수상현황(사용자 제공 xlsx 수상 경력 시트, 2026-10-23). holder=선수/팀/감독. 선수번호 연결은 시즌·이름 일치로 한 것이며 link가 "선수번호 미연결"이면 원문만 보존.\n')
    w = csv.DictWriter(f, fieldnames=list(rows[0]), lineterminator='\n')
    w.writeheader(); w.writerows(rows)
print(f'awards.csv: {len(rows)}행')
